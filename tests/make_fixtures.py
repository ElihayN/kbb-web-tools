# -*- coding: utf-8 -*-
"""Builds synthetic fixture files to exercise each module without needing
Elihay's real (sensitive) business documents.
"""
import os
import datetime as dt
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
FIX = os.path.join(HERE, "fixtures")
os.makedirs(FIX, exist_ok=True)


def make_weekly_report():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    headers = {1: "Cso no.", 2: "Status", 8: "Prod Name", 20: "Arrival Date",
               26: "Customer Name", 27: "Company", 28: "House Phone", 29: "Celular Phone"}
    for col, h in headers.items():
        ws.cell(row=1, column=col, value=h)

    today = dt.date.today()

    def d(days_ago):
        return (today - dt.timedelta(days=days_ago)).strftime("%d/%m/%y")

    rows = [
        # cso, status,             prod,        arrival(days ago), customer,      company, home,        mobile
        ("CSO1001", "Missing Spare Part", "iPhone 14",  2,  "דנה כהן",     "0",       "0521234567", "0541234567"),
        ("CSO1002", "Missing Spare Part", "iPhone 13",  10, "יוסי לוי",    "חברת בע\"מ", "0525556677", ""),
        ("CSO1003", "Missing Product",    "MacBook Air", 1,  "מירי אזולאי", "0",       "",           "0509998877"),
        ("CSO1004", "Missing Product",    "iPad Pro",    9,  "אבי שרון",    "0",       "",           ""),
        ("CSO1005", "Closed",             "iPhone 12",   30, "לא רלוונטי",  "0",       "0500000000", ""),
        ("CSO1006", "Missing Spare Part", "iPhone SE",   0,  "רות ברק",     "0",       "035556677",  "0"),
    ]
    r = 2
    for cso, status, prod, days_ago, cust, company, home, mobile in rows:
        ws.cell(row=r, column=1, value=cso)
        ws.cell(row=r, column=2, value=status)
        ws.cell(row=r, column=8, value=prod)
        ws.cell(row=r, column=20, value=d(days_ago))
        ws.cell(row=r, column=26, value=cust)
        ws.cell(row=r, column=27, value=company)
        ws.cell(row=r, column=28, value=home)
        ws.cell(row=r, column=29, value=mobile)
        r += 1

    path = os.path.join(FIX, "operational_business_report.xlsx")
    wb.save(path)
    print("wrote", path)


if __name__ == "__main__":
    make_weekly_report()
