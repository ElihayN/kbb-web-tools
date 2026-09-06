# -*- coding: utf-8 -*-
"""
weekly-sms — רשימות SMS ללקוחות הממתינים לחלק/מוצר + עדכון App2U.

Deterministic re-implementation of the `weekly-sms` skill: no AI involved,
just the exact filtering / templating rules that were already nailed down.
"""
import os
import re
import datetime as dt

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from . import office

# ---------------------------------------------------------------- columns --
COL_CSO = 1        # A
COL_STATUS = 2     # B
COL_PROD_NAME = 8  # H
COL_ARRIVAL = 20   # T
COL_CUSTOMER = 26  # Z
COL_COMPANY = 27   # AA
COL_HOME_PHONE = 28   # AB
COL_MOBILE_PHONE = 29  # AC

STATUS_MAP = {
    "Missing Spare Part": "חסר חלק",
    "Missing Product": "חסר מוצר",
}

# 4 output lists, in the fixed order they should appear in the workbook.
LIST_KEYS = [
    ("חסר חלק", "קצר", "חסר חלק 0-6"),
    ("חסר חלק", "מתעכב", "חסר חלק מעל 6"),
    ("חסר מוצר", "קצר", "חסר מוצר 0-6"),
    ("חסר מוצר", "מתעכב", "חסר מוצר מעל 6"),
]

SMS_TEMPLATES = {
    ("חסר חלק", "קצר"): (
        "לקוח יקר\n"
        "רצינו לעדכן כי אנו ממתינים לקבלת חלק/ים להשלמת תיקון מוצרך  "
        "בקריאת שירות מס' ( מס קריאה CSO ) , מיד עם הגעת החלק/ים נשלים "
        "את התיקון ונעדכן בהתאם , לשירותך ולרשותך צוות Dcs"
    ),
    ("חסר חלק", "מתעכב"): (
        "לקוח יקר,\n"
        "רצינו לעדכן כי הטיפול במוצר שלך מתעכב עקב המתנה לחלק הנדרש "
        "לתיקון, אשר טרם התקבל מהיצרן Apple.\n"
        "אנו עוקבים אחר אספקת החלק ונפעל להשלים את התיקון בהקדם עם הגעתו.\n"
        "קריאת שירות מס'  ( מס קריאה CSO )\n"
        "מתנצלים על העיכוב ומודים לך על הסבלנות וההבנה,\n"
        "צוות Dcs"
    ),
    ("חסר מוצר", "קצר"): (
        "לקוח יקר\n"
        "רצינו לעדכן כי אנו ממתינים לקבלת מוצר חדש מאפל עבורך בקריאת "
        "שירות מס' ( מס קריאה CSO ) , מיד עם הגעת המוצר החדש נעדכן "
        "בהודעת SMS בהתאם. לשירותך ולרשותך צוות Dcs"
    ),
    ("חסר מוצר", "מתעכב"): (
        "לקוח יקר,\n"
        "רצינו לעדכן כי הטיפול במוצר שלך מתעכב עקב המתנה למוצר הנדרש "
        "לתיקון, אשר טרם התקבל מהיצרן Apple.\n"
        "אנו עוקבים אחר אספקת המוצר ונפעל להשלים את התיקון בהקדם עם הגעתו.\n"
        "קריאת שירות מס' ( מס קריאה CSO )\n"
        "מתנצלים על העיכוב ומודים לך על הסבלנות וההבנה,\n"
        "צוות Dcs"
    ),
}

PLACEHOLDER = "( מס קריאה CSO )"

HEADER_FILL = PatternFill("solid", fgColor="1F2A44")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial")

OUTPUT_COLUMNS = [
    "מס' קריאה (CSO)", "שם לקוח", "חברה", "טלפון לשליחת SMS",
    "מס' קריאה", "טלפון נייד", "טלפון בית", "תאריך הגעה למעבדה",
    "ימי המתנה", "מוצר", "נוסח ה-SMS המלא",
]


class WeeklySmsError(RuntimeError):
    pass


def _cell_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def parse_arrival_date(value):
    """Column T is documented as text DD/MM/YY, but exported reports
    sometimes carry a real datetime cell instead — handle both."""
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    s = str(value).strip()
    for fmt in ("%d/%m/%y", "%d/%m/%Y", "%d-%m-%y", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def clean_phone(value):
    s = _cell_str(value)
    if not s or s == "0":
        return ""
    # numeric phone columns sometimes come through as floats -> "5xxxxxxxx.0"
    s = re.sub(r"\.0$", "", s)
    s = re.sub(r"\D", "", s)
    return s


def clean_company(value):
    s = _cell_str(value)
    if s == "0":
        return ""
    return s


def process_weekly_report(input_path: str, work_dir: str, today: dt.date | None = None):
    """Run the full weekly-sms pipeline.

    Returns a dict: {
        'sms_path': ..., 'app2u_path': ...,
        'counts': {list_name: n, ...},
        'warnings': [str, ...],
    }
    """
    today = today or dt.date.today()
    os.makedirs(work_dir, exist_ok=True)
    warnings = []

    ext = os.path.splitext(input_path)[1].lower()
    if ext == ".xls":
        input_path = office.xls_to_xlsx(input_path, work_dir)

    wb = openpyxl.load_workbook(input_path, data_only=True)
    ws = wb.worksheets[0]

    buckets = {key: [] for key in [(c, b) for c, b, _ in LIST_KEYS]}
    no_phone_csos = []

    for row in ws.iter_rows(min_row=2):
        status_raw = _cell_str(row[COL_STATUS - 1].value)
        category = STATUS_MAP.get(status_raw)
        if not category:
            continue

        cso = _cell_str(row[COL_CSO - 1].value)
        if not cso:
            continue

        arrival = parse_arrival_date(row[COL_ARRIVAL - 1].value)
        if arrival is None:
            warnings.append(f"קריאה {cso}: לא ניתן היה לפרש את תאריך ההגעה — הקריאה דולגה.")
            continue

        waiting_days = (today - arrival).days
        bucket = "קצר" if waiting_days <= 6 else "מתעכב"

        mobile = clean_phone(row[COL_MOBILE_PHONE - 1].value)
        home = clean_phone(row[COL_HOME_PHONE - 1].value)
        sms_phone = mobile or home
        if not sms_phone:
            no_phone_csos.append(cso)

        customer = _cell_str(row[COL_CUSTOMER - 1].value)
        company = clean_company(row[COL_COMPANY - 1].value)
        prod_name = _cell_str(row[COL_PROD_NAME - 1].value)

        sms_text = SMS_TEMPLATES[(category, bucket)].replace(PLACEHOLDER, cso)

        buckets[(category, bucket)].append({
            "cso": cso,
            "customer": customer,
            "company": company,
            "sms_phone": sms_phone,
            "mobile": mobile,
            "home": home,
            "arrival": arrival,
            "waiting_days": waiting_days,
            "prod_name": prod_name,
            "sms_text": sms_text,
        })

    for key in buckets:
        buckets[key].sort(key=lambda r: r["arrival"])  # oldest first (most urgent)

    if no_phone_csos:
        warnings.append(
            "קריאות רלוונטיות ללא שום מספר טלפון (לא נייד ולא בית): "
            + ", ".join(no_phone_csos)
        )

    sms_path = _build_sms_workbook(buckets, today, work_dir)
    sms_path = office.recalc_xlsx(sms_path, work_dir)  # bake formula values

    app2u_xlsx = _build_app2u_workbook(buckets, today, work_dir)
    app2u_path = office.xlsx_to_legacy_xls(app2u_xlsx, work_dir)

    counts = {name: len(buckets[(c, b)]) for c, b, name in LIST_KEYS}

    return {
        "sms_path": sms_path,
        "app2u_path": app2u_path,
        "counts": counts,
        "warnings": warnings,
    }


def _style_header(ws, ncols):
    ws.sheet_view.rightToLeft = True
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 26


def _autosize(ws, ncols, widths=None):
    for col in range(1, ncols + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = (widths or {}).get(col, 18)


def _build_sms_workbook(buckets, today, work_dir):
    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = "סיכום"
    summary.sheet_view.rightToLeft = True
    summary["A1"] = "תאריך הפקה"
    summary["B1"] = today.strftime("%d/%m/%Y")
    summary["A1"].font = Font(name="Arial", bold=True)
    summary["A3"] = "רשימה"
    summary["B3"] = "מס' קריאות"
    _style_header(summary, 2)
    r = 4
    for _, _, name in LIST_KEYS:
        summary.cell(row=r, column=1, value=name).font = BODY_FONT
        summary.cell(row=r, column=2, value=f"=COUNTA('{name}'!A2:A100000)").font = BODY_FONT
        r += 1
    _autosize(summary, 2, {1: 26, 2: 14})

    for category, bucket, name in LIST_KEYS:
        ws = wb.create_sheet(title=name)
        ws.sheet_view.rightToLeft = True
        for c, title in enumerate(OUTPUT_COLUMNS, start=1):
            ws.cell(row=1, column=c, value=title)
        _style_header(ws, len(OUTPUT_COLUMNS))

        rows = buckets[(category, bucket)]
        for i, item in enumerate(rows, start=2):
            ws.cell(row=i, column=1, value=item["cso"])
            ws.cell(row=i, column=2, value=item["customer"])
            ws.cell(row=i, column=3, value=item["company"])
            ws.cell(row=i, column=4, value=item["sms_phone"])
            ws.cell(row=i, column=5, value=f"=A{i}")
            ws.cell(row=i, column=6, value=item["mobile"])
            ws.cell(row=i, column=7, value=item["home"])
            ws.cell(row=i, column=8, value=item["arrival"])
            ws.cell(row=i, column=8).number_format = "DD/MM/YYYY"
            ws.cell(row=i, column=9, value=f"=TODAY()-H{i}")
            ws.cell(row=i, column=10, value=item["prod_name"])
            ws.cell(row=i, column=11, value=item["sms_text"])
            ws.cell(row=i, column=11).alignment = Alignment(wrap_text=True, vertical="top")
            for c in range(1, len(OUTPUT_COLUMNS) + 1):
                ws.cell(row=i, column=c).font = BODY_FONT

        widths = {1: 12, 2: 20, 3: 18, 4: 14, 5: 12, 6: 14, 7: 14, 8: 16, 9: 12, 10: 20, 11: 60}
        _autosize(ws, len(OUTPUT_COLUMNS), widths)
        ws.freeze_panes = "A2"

    # openpyxl never writes a cached result for formulas — force LibreOffice
    # (and Excel) to fully recompute on next open, otherwise some viewers
    # show 0/blank until a human re-saves the file.
    wb.calculation.fullCalcOnLoad = True

    path = os.path.join(work_dir, "רשימות_SMS.xlsx")
    wb.save(path)
    return path


def _build_app2u_workbook(buckets, today, work_dir):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "App2U"
    headers = ["cso", "c", "date", "name", "e", "subject", "body", "comtype"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)

    r = 2
    for key in [(c, b) for c, b, _ in LIST_KEYS]:
        for item in buckets[key]:
            ws.cell(row=r, column=1, value=item["cso"])
            ws.cell(row=r, column=2, value="")
            ws.cell(row=r, column=3, value=today.strftime("%d/%m/%Y"))
            ws.cell(row=r, column=4, value="מערכת")
            ws.cell(row=r, column=5, value="")
            ws.cell(row=r, column=6, value="משלוח חלק / מוצר ")
            ws.cell(row=r, column=7, value=item["sms_text"])
            ws.cell(row=r, column=8, value="SMS")
            r += 1

    path = os.path.join(work_dir, "App2U_Upload.xlsx")
    wb.save(path)
    return path
